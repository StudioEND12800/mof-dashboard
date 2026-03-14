"""
fix_duplicates.py – Supprime les doublons Nom/Prénom inversés et les vrais doublons de MOF_France_FUSIONNÉ.xlsx
Stratégie :
  - Clé canonique = (sorted(Nom_norm, Prénom_norm), Métier_norm)
  - Pour chaque groupe >1 ligne : garder la ligne la plus complète,
    avec priorité source SNMOF_export > SNMOF_scrape > autres
  - Réécrire le fichier Excel proprement
"""

import pandas as pd
import unicodedata
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

EXCEL_PATH = Path("/Users/nicolas/projets_mof/output/MOF_France_FUSIONNÉ.xlsx")

# Priorité des sources (plus bas = meilleur)
SOURCE_PRIORITY = {
    "SNMOF_export":   0,
    "SNMOF_scrape":   1,
    "mof_boulangers": 2,
    "mof_coiffure":   3,
    "portail_chocola":4,
    "portail-du-choc":5,
    "wikipedia-fr":   6,
}


def norm(s):
    if not s or (isinstance(s, float)):
        return ""
    s = str(s).strip().upper()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def canon_key(row):
    n = norm(row["Nom"])
    p = norm(row["Prénom"])
    m = norm(str(row["Métier"])) if pd.notna(row["Métier"]) else ""
    return (tuple(sorted([n, p])), m)


def source_rank(src):
    if not isinstance(src, str):
        return 99
    for key, rank in SOURCE_PRIORITY.items():
        if key in src:
            return rank
    return 50


def count_non_null(row):
    """Compte le nombre de colonnes non-vides pour scorer la complétude."""
    return sum(1 for v in row if pd.notna(v) and str(v).strip() != "")


def merge_rows(group: pd.DataFrame) -> pd.Series:
    """
    Fusionne un groupe de doublons en gardant la meilleure ligne
    (priorité source + complétude) et en remplissant les champs vides
    avec les infos des autres lignes.
    """
    # Trier : source prioritaire d'abord, puis la plus complète
    group = group.copy()
    group["_src_rank"] = group["Source"].apply(source_rank)
    group["_completeness"] = group.apply(count_non_null, axis=1)
    group = group.sort_values(["_src_rank", "_completeness"], ascending=[True, False])

    best = group.iloc[0].copy()

    # Enrichir les champs vides avec les autres lignes
    fill_cols = ["Nom", "Prénom", "Nom complet", "Département", "Région",
                 "Num_Dept", "Ville", "Classe", "Entreprise", "Fiabilité"]
    for col in fill_cols:
        if col not in group.columns:
            continue
        if pd.isna(best[col]) or str(best[col]).strip() == "":
            for _, other in group.iloc[1:].iterrows():
                if pd.notna(other[col]) and str(other[col]).strip() != "":
                    best[col] = other[col]
                    break

    # Garder le Nom complet de la meilleure ligne (SNMOF_export a toujours l'ordre correct)
    # On le recompose uniquement si le champ est vide
    if pd.isna(best.get("Nom complet")) or str(best.get("Nom complet", "")).strip() == "":
        nom_str = str(best["Nom"]).strip() if pd.notna(best["Nom"]) else ""
        prenom_str = str(best["Prénom"]).strip() if pd.notna(best["Prénom"]) else ""
        if nom_str and prenom_str:
            best["Nom complet"] = f"{nom_str.upper()} {prenom_str}"

    return best.drop(labels=["_src_rank", "_completeness"], errors="ignore")


def main():
    print(f"Chargement : {EXCEL_PATH}")
    df = pd.read_excel(EXCEL_PATH)
    n_orig = len(df)
    print(f"  {n_orig} lignes chargées")

    # Calcul clé canonique
    df["_ckey"] = df.apply(canon_key, axis=1)

    # Séparer unicité et doublons
    singles_mask = ~df.duplicated(subset=["_ckey"], keep=False)
    singles = df[singles_mask].copy()
    dupes = df[~singles_mask].copy()

    print(f"  Lignes sans doublon : {len(singles)}")
    print(f"  Lignes avec doublon (à traiter) : {len(dupes)}")

    # Traiter chaque groupe
    fixed_rows = []
    inv_count = 0
    dup_count = 0

    for key, grp in dupes.groupby("_ckey"):
        noms   = grp["Nom"].apply(norm).tolist()
        prenoms = grp["Prénom"].apply(norm).tolist()
        vals = set(zip(noms, prenoms))
        if len(vals) > 1:
            inv_count += 1
        else:
            dup_count += 1
        fixed_rows.append(merge_rows(grp))

    print(f"  Groupes inversions corrigés : {inv_count}")
    print(f"  Groupes vrais doublons dédupliqués : {dup_count}")

    # Recomposer le DataFrame final
    df_fixed = pd.concat([singles, pd.DataFrame(fixed_rows)], ignore_index=True)

    # Nettoyer colonnes temporaires
    df_fixed = df_fixed.drop(columns=["_ckey"], errors="ignore")

    # Trier
    df_fixed = df_fixed.sort_values(
        ["Num_Dept", "Promotion", "Nom"],
        key=lambda col: col.fillna("").astype(str)
    ).reset_index(drop=True)

    n_final = len(df_fixed)
    print(f"\n  Avant : {n_orig} | Après : {n_final} | Supprimés : {n_orig - n_final}")

    # ----- Sauvegarde Excel -----
    print(f"\nSauvegarde → {EXCEL_PATH}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MOFs"

    header_fill = PatternFill(start_color="1A2E5A", end_color="1A2E5A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # Headers
    cols = list(df_fixed.columns)
    for ci, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data
    for ri, row_data in enumerate(df_fixed.itertuples(index=False), 2):
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=ri, column=ci, value=val if pd.notna(val) else None)

    # Auto-width
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)

    # Freeze + autofilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Feuille stats métiers
    ws2 = wb.create_sheet("Stats Métiers")
    stats_m = df_fixed.groupby("Métier").size().sort_values(ascending=False).reset_index(name="Nb")
    ws2.append(["Métier", "Nb"])
    for _, r in stats_m.iterrows():
        ws2.append([r["Métier"], r["Nb"]])

    # Feuille stats département
    ws3 = wb.create_sheet("Stats Départements")
    stats_d = df_fixed.groupby(["Num_Dept", "Département", "Région"]).size().sort_values(ascending=False).reset_index(name="Nb")
    ws3.append(["Num_Dept", "Département", "Région", "Nb"])
    for _, r in stats_d.iterrows():
        ws3.append([r["Num_Dept"], r["Département"], r["Région"], r["Nb"]])

    wb.save(EXCEL_PATH)
    print(f"  ✓ Fichier sauvegardé ({n_final} MOFs)")

    # Stats rapides
    print("\n=== Vérification finale ===")
    df2 = pd.read_excel(EXCEL_PATH)
    print(f"  Lignes Excel : {len(df2)}")
    print(f"  Métiers uniques : {df2['Métier'].nunique()}")
    print(f"  Départements renseignés : {df2['Département'].notna().sum()}")


if __name__ == "__main__":
    main()
